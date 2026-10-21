# -*- coding: utf-8 -*-
"""生成示例用的小样片 assets/demo.pptx（4 页），给两个示例 plan 当靶子。

为什么要这个脚本：plan 里的 spid 是**形状 id**，跟具体文件绑定。
没有配套 pptx，"示例"就只是纸上谈兵 —— 文档里写的命令一跑就报错。
这个脚本生成的形状 id 是确定的（顺序建），两份 plan 直接对着它写。

用法：python scripts/make_demo_deck.py
产出：assets/demo.pptx（960×540，4 页：封面 / 内容 / 内容 / 满版收尾）
"""
import io
import math
import os

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(os.path.dirname(HERE), "assets")
os.makedirs(ASSETS, exist_ok=True)
OUT = os.path.join(ASSETS, "demo.pptx")

W, H = 960.0, 540.0
BG = (14, 18, 17)
GOLD = "E3C08A"
INK = "F2EDE2"


def emu(v):
    return Emu(int(v * 12700))


def C(h):
    return RGBColor.from_string(h)


def _rect(s, x, y, w, h, color, alpha=0, name=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    sp.shadow.inherit = False
    st = sp._element.find(qn("p:style"))
    if st is not None:
        sp._element.remove(st)
    sp.fill.solid()
    sp.fill.fore_color.rgb = C(color)
    sp.line.fill.background()
    if alpha:
        sf = sp._element.spPr.find(qn("a:solidFill"))
        clr = sf.find(qn("a:srgbClr"))
        al = etree.SubElement(clr, qn("a:alpha"))
        al.set("val", str(int((100 - alpha) * 1000)))
    if name:
        sp.name = name
    return sp


def _ring(s, cx, cy, d, color=GOLD, w=1.2, name=None):
    sp = s.shapes.add_shape(MSO_SHAPE.OVAL, emu(cx - d / 2), emu(cy - d / 2), emu(d), emu(d))
    sp.shadow.inherit = False
    st = sp._element.find(qn("p:style"))
    if st is not None:
        sp._element.remove(st)
    sp.fill.background()
    sp.line.color.rgb = C(color)
    sp.line.width = Pt(w)
    if name:
        sp.name = name
    return sp


def _arc(s, cx, cy, d, a0, a1, color=GOLD, w=2.0, name="arc"):
    sp = s.shapes.add_shape(MSO_SHAPE.ARC, emu(cx - d / 2), emu(cy - d / 2), emu(d), emu(d))
    sp.shadow.inherit = False
    st = sp._element.find(qn("p:style"))
    if st is not None:
        sp._element.remove(st)
    sp.fill.background()
    sp.line.color.rgb = C(color)
    sp.line.width = Pt(w)
    geom = sp._element.spPr.find(qn("a:prstGeom"))
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for nm, val in (("adj1", a0), ("adj2", a1)):
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", nm)
        gd.set("fmla", f"val {int(val * 60000)}")
    sp.name = name
    return sp


def _txt(s, x, y, w, h, t, size, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bold=False, name=None, font="Microsoft YaHei"):
    tb = s.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = t
    p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.name = font
        run.font.bold = bold
        run.font.color.rgb = C(color)
    if name:
        tb.name = name
    return tb


def _disc_img():
    """圆窗用的柔和渐变图（不依赖 numpy，PIL 逐点算）。"""
    N = 480
    R = N / 2
    px = []
    for y in range(N):
        for x in range(N):
            t = min(1.0, math.hypot(x - R, y - R) / R)
            a = (1 - t) ** 1.6
            px.append((int(196 * a + 24), int(170 * a + 26), int(126 * a + 28)))
    im = Image.new("RGB", (N, N))
    im.putdata(px)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    buf.seek(0)
    return buf



def _pic_round(s, cx, cy, d, name="coreImg"):
    """图片裁成圆形（prstGeom=ellipse）—— 示例里必须裁，否则圆窗是个方图。"""
    pic = s.shapes.add_picture(IMG, emu(cx - d / 2), emu(cy - d / 2), emu(d), emu(d))
    pic.name = name
    spPr = pic._element.spPr
    geom = spPr.find(qn("a:prstGeom"))
    if geom is not None:
        geom.set("prst", "ellipse")
        if geom.find(qn("a:avLst")) is None:
            etree.SubElement(geom, qn("a:avLst"))
    return pic


IMG = _disc_img()

prs = Presentation()
prs.slide_width = emu(W)
prs.slide_height = emu(H)
BLANK = prs.slide_layouts[6]

# ── P1 封面：圆在右 ────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, "0E1211", name="bg")
_pic_round(s, 792, 246, 488)
_ring(s, 792, 246, 488, GOLD, 1.4, name="coreRing")
_arc(s, 792, 246, 520, -40, 4, GOLD, 2.0, name="arcA0")
_txt(s, 58, 150, 440, 120, "示例小样片", 46, INK, bold=True, name="hd")
_txt(s, 58, 286, 440, 20, "DEMO DECK  /  4 PAGES", 11, GOLD, name="eyebrow")
_rect(s, 58, 320, 44, 2.4, GOLD, name="rule")

# ── P2 内容：圆在左 ────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, "0E1211", name="bg")
_pic_round(s, 168, 246, 420)
_ring(s, 168, 246, 420, GOLD, 1.4, name="coreRing")
_arc(s, 168, 246, 452, 140, 184, GOLD, 2.0, name="arcA0")
_txt(s, 432, 96, 470, 40, "三行文字，逐条入场", 26, INK, bold=True, name="hd")
_rect(s, 432, 150, 40, 2.4, GOLD, name="rule")
for i, t in enumerate(["第一条：cinematic 入场", "第二条：错峰 80ms", "第三条：最后落定"]):
    _txt(s, 432, 190 + i * 40, 470, 18, t, 12, INK, name=f"row{i}")

# ── P3 内容：圆在右（走位 + 尺寸变化）──────────────────────────
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, "0E1211", name="bg")
_pic_round(s, 772, 246, 360)
_ring(s, 772, 246, 360, GOLD, 1.4, name="coreRing")
_arc(s, 772, 246, 392, -60, -16, GOLD, 2.0, name="arcA0")
_txt(s, 58, 96, 470, 40, "圆窗换了位置和大小", 26, INK, bold=True, name="hd")
_rect(s, 58, 150, 40, 2.4, GOLD, name="rule")
_txt(s, 58, 190, 470, 60, "翻页时圆从上一页的位置滑过来，直径同步推拉。", 12, INK, name="row0")

# ── P4 满版收尾 ────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, "0E1211", name="bg")
_pic_round(s, 480, 270, 1080)
_ring(s, 480, 270, 1080, GOLD, 1.2, name="coreRing")
_txt(s, 130, 210, 700, 60, "谢谢观看", 40, INK, bold=True, name="hd")
_txt(s, 130, 286, 560, 20, "这个文件只是为了跑通示例 plan，不用改它。", 11, GOLD, name="lead")

prs.save(OUT)
print("->", OUT)
